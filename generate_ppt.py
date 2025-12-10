#!/usr/bin/env python3
"""
Enhanced PPTX generator for reproduction results.

Reads JSON or CSV result files produced by `scripts/eval.py` (or accepts manual fields)
and writes a structured PowerPoint summary with slides for: Title, Overview, Data,
Training / Checkpoint, Quantitative Results, Visualizations (placeholders), and Next Steps.

Requires `python-pptx` (pip install python-pptx).

Usage examples:
  python3 scripts/generate_ppt.py --json output/eval_224.json --out output/repro_report.pptx
  python3 scripts/generate_ppt.py --csv output/eval_224.csv --out output/repro_report.pptx
"""
import argparse
import json
import csv
import os
from pptx import Presentation
from pptx.util import Inches, Pt


def add_bullets(slide, lines, font_size=18):
    tx = slide.shapes.placeholders[1].text_frame
    tx.clear()
    for i, l in enumerate(lines):
        p = tx.add_paragraph() if i > 0 else tx.paragraphs[0]
        p.text = l
        p.font.size = Pt(font_size)


def make_title(prs, title, subtitle=None):
    slide = prs.slides.add_slide(prs.slide_layouts[0])
    slide.shapes.title.text = title
    if subtitle:
        try:
            slide.placeholders[1].text = subtitle
        except Exception:
            pass


def make_slide(prs, title, lines, font_size=18):
    slide = prs.slides.add_slide(prs.slide_layouts[1])
    slide.shapes.title.text = title
    add_bullets(slide, lines, font_size=font_size)


def read_json(path):
    with open(path, 'r') as f:
        return json.load(f)


def read_csv(path):
    d = {}
    with open(path, 'r', newline='') as f:
        r = csv.reader(f)
        for row in r:
            if not row:
                continue
            key = row[0]
            val = row[1] if len(row) > 1 else ''
            # try to parse JSON value in cell
            try:
                parsed = json.loads(val)
                d[key] = parsed
            except Exception:
                d[key] = val
    return d


def build_deck(prs, meta):
    # Title
    title = meta.get('title', 'LookHere Reproduction — Results')
    subtitle = meta.get('subtitle', 'Auto-generated')
    make_title(prs, title, subtitle)

    # Overview
    overview = [
        f"Repository: {meta.get('repo','(unknown)')}",
        f"Model: {meta.get('model','LookHere / ViT')}",
        f"Checkpoint: {meta.get('checkpoint','(none)')}",
        f"Eval task: {meta.get('task','hr/classify/all')}",
    ]
    make_slide(prs, 'Overview', overview)

    # Data
    data_lines = [
        f"Dataset: {meta.get('dataset','ImageNet')}",
        f"Val dir: {meta.get('val_dir','(unspecified)')}",
        f"HR eval images: {meta.get('total', meta.get('hr_total', 'N/A'))}"
    ]
    make_slide(prs, 'Data', data_lines)

    # Training / Checkpoint
    train_lines = [
        f"Epochs: {meta.get('epochs','150')}",
        f"Init: {meta.get('init','author checkpoint')}",
        f"Image size (train): {meta.get('img_size','224')}",
        f"Patch size: {meta.get('patch_size','16')}",
    ]
    make_slide(prs, 'Training / Checkpoint', train_lines)

    # Quantitative results
    q = []
    q.append(f"Total samples: {meta.get('total','N/A')}")
    if 'top1' in meta:
        q.append(f"Top-1: {meta.get('top1')}")
    if 'top5' in meta:
        q.append(f"Top-5: {meta.get('top5')}")
    if 'ece' in meta:
        q.append(f"ECE: {meta.get('ece')}")
    if 'fgsm' in meta and isinstance(meta['fgsm'], dict):
        q.append('FGSM results:')
        for eps, acc in meta['fgsm'].items():
            q.append(f"  eps={eps}: {acc}")
    make_slide(prs, 'Quantitative Results', q)

    # Visualizations placeholders
    viz_lines = [
        'Attention maps (examples): add figures',
        'HR vs baseline image examples: add figures',
        'ECE reliability diagram: add figure',
    ]
    make_slide(prs, 'Visualizations (placeholders)', viz_lines)

    # Next steps / conclusions
    ns = [
        'Mapping: remaining unknown HR images moved to unknown_class',
        'Option: run 5-epoch finetune at 384px if HR results are poor',
        'Proposed: remap unknowns or accept 4,989 subset for final report',
    ]
    make_slide(prs, 'Next Steps', ns)


def main():
    p = argparse.ArgumentParser()
    p.add_argument('--json', type=str, default=None, help='results JSON file')
    p.add_argument('--csv', type=str, default=None, help='results CSV file')
    p.add_argument('--checkpoint', type=str, default=None, help='checkpoint path (for slide)')
    p.add_argument('--val-dir', type=str, default=None, help='validation directory used')
    p.add_argument('--repo', type=str, default=None, help='repository name or path')
    p.add_argument('--out', type=str, required=True, help='output PPTX path')
    args = p.parse_args()

    meta = {}
    if args.repo:
        meta['repo'] = args.repo
    if args.checkpoint:
        meta['checkpoint'] = args.checkpoint
    if args.val_dir:
        meta['val_dir'] = args.val_dir

    # read results
    if args.json and os.path.isfile(args.json):
        try:
            j = read_json(args.json)
            meta.update(j)
        except Exception:
            print('Warning: failed to read JSON results file')
    elif args.csv and os.path.isfile(args.csv):
        try:
            c = read_csv(args.csv)
            meta.update(c)
        except Exception:
            print('Warning: failed to read CSV results file')

    prs = Presentation()
    build_deck(prs, meta)
    out_dir = os.path.dirname(args.out)
    if out_dir and not os.path.isdir(out_dir):
        os.makedirs(out_dir, exist_ok=True)
    prs.save(args.out)
    print('Wrote PPTX to', args.out)


if __name__ == '__main__':
    main()

